#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Catálogo de atribuições de número de manifestantes na imprensa — 8 de Janeiro.

Instrumento SEPARADO do codebook. Não alimenta o process tracing nem a cronologia
interativa: serve como aparato de CONFERÊNCIA e como material para a análise de
como os veículos disputam discursivamente o processo político pela via do número.

O codebook mantém apenas os números do documento legal (CPMI e anexos oficiais).

Gera: data/processed/CATALOGO-IMPRENSA-MULTIDAO-8JAN.xlsx

Uso:  python scripts/gerar_catalogo_imprensa.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "data/processed/CATALOGO-IMPRENSA-MULTIDAO-8JAN.xlsx"
ACESSO = "16 ago. 2026"

# ---------------------------------------------------------------------------
# Aba 1 — Atribuições da imprensa
# Colunas pedidas: número · link · citação ABNT. As demais existem porque sem
# elas o número não se sustenta: quem atribuiu, a quê se refere e como foi
# verificado é o que distingue catálogo de coletânea.
# ---------------------------------------------------------------------------
COL_IMPRENSA = [
    ("Veículo", 26), ("Número de manifestantes", 30), ("Referente do número", 34),
    ("Atribuição declarada pelo veículo", 30), ("Data de publicação", 16),
    ("Natureza", 18), ("Link (endereço web)", 62), ("Citação (ABNT)", 78),
    ("Verificação", 30),
]

IMPRENSA = [
    # --- verificados por leitura direta da página -------------------------
    ("Agência Brasil (EBC)",
     "cerca de 4 mil",
     "Praça dos Três Poderes, às 15h26 do dia 08/01",
     '"Policiais estimam" — sem órgão nomeado',
     "jan. 2024", "Retrospectiva",
     "https://agenciabrasil.ebc.com.br/politica/noticia/2024-01/relembre-o-passo-passo-da-tentativa-de-golpe-no-81",
     "RELEMBRE o passo a passo da tentativa de golpe no 8/1. **Agência Brasil**, Brasília, jan. 2024. "
     "Disponível em: https://agenciabrasil.ebc.com.br/politica/noticia/2024-01/relembre-o-passo-passo-da-tentativa-de-golpe-no-81. "
     "Acesso em: " + ACESSO + ".",
     "VERIFICADO — página lida"),

    ("Agência Brasil (EBC)",
     "300 (05/01) e 5.500 (07/01)",
     "Acampamento do Setor Militar Urbano",
     "CPMI, para os 5.500; reportagem, para os 300",
     "jan. 2024", "Retrospectiva",
     "https://agenciabrasil.ebc.com.br/politica/noticia/2024-01/relembre-o-passo-passo-da-tentativa-de-golpe-no-81",
     "RELEMBRE o passo a passo da tentativa de golpe no 8/1. **Agência Brasil**, Brasília, jan. 2024. "
     "Disponível em: https://agenciabrasil.ebc.com.br/politica/noticia/2024-01/relembre-o-passo-passo-da-tentativa-de-golpe-no-81. "
     "Acesso em: " + ACESSO + ".",
     "VERIFICADO — confere com CPMI p. 576"),

    ("Agência Brasil (EBC)",
     "NENHUM — apenas \"milhares de apoiadores\"",
     "Dia 08/01, sem quantificação",
     "—",
     "jan. 2026", "Retrospectiva",
     "https://agenciabrasil.ebc.com.br/politica/noticia/2026-01/cronologia-da-tentativa-de-golpe-atentados-bloqueios-e-acampamentos",
     "CRONOLOGIA da tentativa de golpe: atentados, bloqueios e acampamentos. **Agência Brasil**, Brasília, jan. 2026. "
     "Disponível em: https://agenciabrasil.ebc.com.br/politica/noticia/2026-01/cronologia-da-tentativa-de-golpe-atentados-bloqueios-e-acampamentos. "
     "Acesso em: " + ACESSO + ".",
     "VERIFICADO — o próprio veículo abandona a cifra"),

    ("Poder360",
     "NENHUM para a Praça. Traz \"cerca de 100 pessoas\" (13h) e \"cerca de 80 ônibus\"",
     "Frente do Congresso às 13h; ônibus chegados em 07/01",
     "Reportagem própria",
     "15 jan. 2023", "Contemporânea",
     "https://poder360.com.br/poder-governo/governo/invasao-aos-tres-poderes-completa-uma-semana-relembre",
     "INVASÃO aos Três Poderes completa uma semana; relembre. **Poder360**, Brasília, 15 jan. 2023. "
     "Disponível em: https://poder360.com.br/poder-governo/governo/invasao-aos-tres-poderes-completa-uma-semana-relembre. "
     "Acesso em: " + ACESSO + ".",
     "VERIFICADO — cronologia contemporânea SEM a cifra de 4 mil"),

    ("CNN Brasil",
     "NENHUM de participantes",
     "Retrospectiva de três anos; quantifica só condenações",
     "—",
     "8 jan. 2026", "Retrospectiva",
     "https://www.cnnbrasil.com.br/politica/8-1-relembre-o-dia-dos-ataques-aos-tres-poderes-ocorridos-ha-tres-anos/",
     "8/1: relembre o dia dos ataques aos Três Poderes, ocorridos há três anos. **CNN Brasil**, São Paulo, 8 jan. 2026. "
     "Disponível em: https://www.cnnbrasil.com.br/politica/8-1-relembre-o-dia-dos-ataques-aos-tres-poderes-ocorridos-ha-tres-anos/. "
     "Acesso em: " + ACESSO + ".",
     "VERIFICADO — página lida"),

    ("Wikipédia lusófona",
     "cerca de 4 mil (duas passagens)",
     "(a) os que SAÍRAM DO QG; (b) o acampamento que \"aumentou para mais de quatro mil\"",
     "Referências não resolvidas no corpo do verbete",
     "consultado 2026", "Terciária",
     "https://pt.wikipedia.org/wiki/Ataques_de_8_de_janeiro_em_Bras%C3%ADlia",
     "ATAQUES de 8 de janeiro em Brasília. In: **Wikipédia**: a enciclopédia livre. [S. l.], 2026. "
     "Disponível em: https://pt.wikipedia.org/wiki/Ataques_de_8_de_janeiro_em_Brasília. Acesso em: " + ACESSO + ".",
     "VERIFICADO — NÃO citável; registrado por evidenciar a migração do número"),

    ("Wikipédia lusófona",
     "3.900 passageiros em 105 veículos",
     "Transporte fretado, alertas da Abin",
     "Abin, via WhatsApp",
     "consultado 2026", "Terciária",
     "https://pt.wikipedia.org/wiki/Ataques_de_8_de_janeiro_em_Bras%C3%ADlia",
     "ATAQUES de 8 de janeiro em Brasília. In: **Wikipédia**: a enciclopédia livre. [S. l.], 2026. "
     "Disponível em: https://pt.wikipedia.org/wiki/Ataques_de_8_de_janeiro_em_Brasília. Acesso em: " + ACESSO + ".",
     "CONFERE com CPMI p. 374 e 637 — citar a CPMI, não o verbete"),

    ("OBCC / UFSM (acadêmico)",
     "NENHUM — \"centenas de pessoas acampadas\"",
     "Acampamento",
     "—",
     "2023-2024", "Acadêmica",
     "https://www.ufsm.br/projetos/institucional/observatorio-crise/ataque-a-democracia-brasileira-a-invasao-aos-tres-poderes-em-8-de-janeiro-de-2023",
     "OBSERVATÓRIO da Crise. Ataque à democracia brasileira: a invasão aos Três Poderes em 8 de janeiro de 2023. "
     "**UFSM**, Santa Maria, 2023. Disponível em: https://www.ufsm.br/projetos/institucional/observatorio-crise/"
     "ataque-a-democracia-brasileira-a-invasao-aos-tres-poderes-em-8-de-janeiro-de-2023. Acesso em: " + ACESSO + ".",
     "VERIFICADO — página lida"),

    # --- bloqueados ao agente: exigem coleta pelo pesquisador -------------
    ("Folha de S.Paulo",
     "A COLETAR",
     "A COLETAR",
     "A COLETAR",
     "8-9 jan. 2023", "Contemporânea",
     "https://www1.folha.uol.com.br/  (busca no acervo)",
     "[preencher após coleta — modelo: SOBRENOME, Nome. Título da matéria. **Folha de S.Paulo**, "
     "São Paulo, 8 jan. 2023. Disponível em: URL. Acesso em: " + ACESSO + ".]",
     "BLOQUEADO — uol.com.br recusa o agente; ver aba Protocolo_Coleta"),

    ("O Globo",
     "A COLETAR",
     "A COLETAR",
     "A COLETAR",
     "8-9 jan. 2023", "Contemporânea",
     "https://oglobo.globo.com/  (busca no acervo)",
     "[preencher após coleta — modelo: SOBRENOME, Nome. Título da matéria. **O Globo**, "
     "Rio de Janeiro, 8 jan. 2023. Disponível em: URL. Acesso em: " + ACESSO + ".]",
     "BLOQUEADO — globo.com recusa o agente; ver aba Protocolo_Coleta"),

    ("O Estado de S. Paulo",
     "A COLETAR",
     "A COLETAR",
     "A COLETAR",
     "8-9 jan. 2023", "Contemporânea",
     "https://www.estadao.com.br/  (busca no acervo)",
     "[preencher após coleta — modelo: SOBRENOME, Nome. Título da matéria. **O Estado de S. Paulo**, "
     "São Paulo, 8 jan. 2023. Disponível em: URL. Acesso em: " + ACESSO + ".]",
     "BLOQUEADO — estadao.com.br recusa o agente; ver aba Protocolo_Coleta"),

    ("G1 (Grupo Globo)",
     "A COLETAR",
     "A COLETAR",
     "A COLETAR",
     "8-9 jan. 2023", "Contemporânea",
     "https://g1.globo.com/  (busca no acervo)",
     "[preencher após coleta — modelo: TÍTULO da matéria. **G1**, Brasília, 8 jan. 2023. "
     "Disponível em: URL. Acesso em: " + ACESSO + ".]",
     "BLOQUEADO — globo.com recusa o agente; ver aba Protocolo_Coleta"),
]

# ---------------------------------------------------------------------------
# Aba 2 — Números do documento legal (os que FICAM no process tracing)
# ---------------------------------------------------------------------------
COL_LEGAL = [
    ("Número", 30), ("Referente", 40), ("Órgão que mediu", 30),
    ("Momento", 20), ("Fonte e página", 26), ("Registro no codebook", 20),
    ("Método", 26),
]

LEGAL = [
    ("300 pessoas", "Acampamento do SMU", "Comando Militar do Planalto", "05/01/2023",
     "CPMI, p. 576", "PT-21", "Contagem de efetivo acampado"),
    ("105 ônibus, ~3.900 passageiros", "Transporte fretado com destino a Brasília",
     "ANTT, difundido pela Abin", "07/01/2023, 12h", "CPMI, p. 374 e 637", "PT-21",
     "Contagem de transporte"),
    ("5.500 pessoas", "Acampamento do SMU", "Comando Militar do Planalto (estimativa declarada)",
     "07/01/2023", "CPMI, p. 576", "PT-21", "Estimativa de efetivo acampado"),
    ("101 veículos", "Ônibus chegados até as 8h20", "Abin", "08/01/2023, 8h20",
     "CPMI, p. 643", "PT-22", "Contagem de transporte"),
    ("3.000 pessoas", "Público na área próxima ao SMU", "SSP/DF, ao grupo CIISP-DF",
     "08/01/2023, 9h", "CPMI, p. 643", "PT-22", "Contagem policial de público"),
    ("\"poucos milhares\"", "O grupo que invadiu, depredou e saqueou",
     "Análise da própria CPMI", "08/01/2023", "CPMI, p. 695", "PT-34",
     "Caracterização analítica"),
    ("~3.000; 1.500 saíram por conta própria", "Acampamento na desocupação",
     "CPMI", "09/01/2023, manhã", "CPMI, p. 687", "PT-52", "Contagem na operação"),
    ("2.151 presos", "Detidos no ato de desocupação", "CPMI", "09/01/2023",
     "CPMI, p. 688", "PT-52", "Contagem de detidos"),
    ("4 mil (NÃO USAR para o 8/1)", "Pico do acampamento em 15/11/2022",
     "CPMI", "15/11/2022", "CPMI, p. 532", "PT-12",
     "Outro evento — origem de confusão frequente"),
    ("4 mil (NÃO CITÁVEL)", "Invocado em depoimento do gen. Augusto Heleno",
     "Depoente investigado", "oitiva na CPMI", "CPMI, p. 1324", "—",
     "Fala de investigado, não achado"),
]

# ---------------------------------------------------------------------------
# Aba 3 — Leitura analítica
# ---------------------------------------------------------------------------
COL_ANALISE = [("Achado", 34), ("Evidência", 76), ("Consequência para a redação", 62)]

ANALISE = [
    ("O mesmo número migra entre três referentes",
     "\"4 mil\" aparece como: (a) público na Praça às 15h26 (Agência Brasil, jan/2024); "
     "(b) os que saíram do QG (Wikipédia); (c) o acampamento que \"aumentou para mais de quatro mil\" "
     "(Wikipédia, outra passagem). Três objetos distintos, um único número.",
     "Número que muda de referente sem mudar de valor foi copiado, não medido. É o indício "
     "mais forte contra a cifra, e é reportável como achado sobre a cobertura."),

    ("A cifra tem fonte única, não convergência",
     "Aparece com atribuição vaga (\"policiais\", sem órgão) em uma retrospectiva da EBC de "
     "jan/2024. Não aparece na cronologia contemporânea do Poder360 (15/01/2023), nem na "
     "cronologia posterior da própria EBC (jan/2026), nem na retrospectiva da CNN (jan/2026).",
     "Não citar como \"estimativa policial\". Se for mencionada, é como objeto — exemplo de "
     "cifra que circulou sem lastro —, nunca como dado."),

    ("O veículo de origem abandonou o número",
     "A EBC publica a cifra em jan/2024 e, na cronologia de jan/2026, escreve apenas "
     "\"milhares de apoiadores\".",
     "Fato reportável: a própria fonte recuou. Sustenta o argumento sobre instabilidade "
     "da quantificação na cobertura."),

    ("A cobertura contemporânea não quantificou a multidão",
     "A cronologia do Poder360 de 15/01/2023 traz ônibus (80) e presos (1.395), mas nenhuma "
     "estimativa de público. Quantifica o que é contável e se abstém do resto.",
     "A ausência é informativa e vale um parágrafo: o número da multidão surge depois, na "
     "retrospectiva, não na cobertura do fato."),

    ("Há literatura acadêmica sobre o enquadramento dos três jornais — obtida e conferida",
     "DRUMMOND, D.; MANGINELLI, F. O ataque de 8 de janeiro nos editoriais dos jornais "
     "brasileiros. SciELO Preprints, 2023. DOI 10.1590/SciELOPreprints.7040. Corpus de 105 "
     "editoriais de 1 a 18/01/2023 (Estadão 36, O Globo 35, Folha 34), com análise temática e "
     "de enquadramento. Achados: antes dos ataques não havia preocupação dos jornais com os "
     "acampamentos; a Folha foi o ÚNICO a não vincular os atacantes a Bolsonaro; O Globo usa "
     "esse enquadramento uma vez, e no Estadão ele é frequente. Em "
     "Fontes/08_Imprensa-enquadramento/.",
     "Sustenta a tese de disputa discursiva — mas NÃO preenche a coluna de números: conferido "
     "que o artigo não traz cifra alguma de manifestantes. É preprint, não revisado por pares; "
     "verificar publicação definitiva antes de citar."),

    ("Duas medições oficiais independentes convergem",
     "ANTT (105 ônibus, ~3.900 passageiros, CPMI p. 374 e 637) e SSP/DF (3.000 de público às "
     "9h, CPMI p. 643). Métodos distintos, mesma ordem de grandeza.",
     "É esta convergência que sustenta a faixa de 3.000 a 5.500 no codebook. A imprensa não "
     "participa dessa ancoragem — participa da disputa sobre ela."),
]

# ---------------------------------------------------------------------------
# Aba 4 — Protocolo de coleta do que está bloqueado
# ---------------------------------------------------------------------------
COL_PROTO = [("Passo", 10), ("Ação", 66), ("Detalhe", 74)]

PROTOCOLO = [
    ("—", "POR QUE ESTA ABA EXISTE",
     "Folha (uol.com.br), O Globo e G1 (globo.com) e Estadão (estadao.com.br) recusam o agente "
     "automatizado — a recusa é do publisher, não falha da busca. Nenhum número desses veículos "
     "foi inserido no catálogo sem leitura direta, porque durante esta pesquisa o sumarizador de "
     "busca chegou a atribuir ao Poder360 uma frase que a página não continha. Em catálogo de "
     "conferência, snippet não verificado é pior que célula vazia."),
    ("1", "Abrir o acervo de cada veículo na data do fato",
     "Folha: www1.folha.uol.com.br — editoria Poder, 8 e 9 jan. 2023. "
     "O Globo: oglobo.globo.com — editoria Política. "
     "Estadão: estadao.com.br — editoria Política. G1: g1.globo.com — Política."),
    ("2", "Buscar os termos que costumam carregar a cifra",
     "\"bolsonaristas invadem\", \"manifestantes\", \"mil pessoas\", \"estimativa\", "
     "\"Praça dos Três Poderes\", \"Esplanada dos Ministérios\"."),
    ("3", "Preferir a cobertura do próprio dia à retrospectiva",
     "A pesquisa mostrou que a cifra de 4 mil NÃO está na cobertura contemporânea e aparece só "
     "em retrospectiva. Datar a matéria é parte do achado — registre a data de publicação."),
    ("4", "Registrar SEMPRE quatro coisas, não só o número",
     "(a) o número; (b) o REFERENTE — Praça? Esplanada? acampamento? os que saíram do QG?; "
     "(c) a ATRIBUIÇÃO declarada — \"segundo a PMDF\", \"segundo a SSP-DF\", \"policiais\", "
     "ou nenhuma; (d) a data. Sem o referente, números de veículos diferentes não são comparáveis."),
    ("5", "Arquivar a página",
     "Salvar em PDF ou HTML em Fontes/08_Arquivamento-web/ e submeter a URL ao "
     "web.archive.org/save/. Matéria jornalística sai do ar; evidência arquivada, não."),
    ("6", "Preencher as linhas marcadas A COLETAR",
     "As quatro linhas bloqueadas da aba Atribuicoes_Imprensa já têm o modelo ABNT pronto — "
     "basta completar autor, título e URL."),
    ("7", "NÃO levar nada disto ao codebook",
     "Por decisão do pesquisador em 16/08/2026, o process tracing mantém apenas os números do "
     "documento legal (aba Numeros_Documento_Legal). Este catálogo é aparato de conferência e "
     "material de análise da cobertura — não entra na cronologia interativa."),
]

FILL_CAB = PatternFill("solid", fgColor="7B2D26")
FILL_ALT = PatternFill("solid", fgColor="F5EFEE")
FILL_BLOQ = PatternFill("solid", fgColor="FDF0D5")
FONT_CAB = Font(bold=True, color="FFFFFF", size=10)
BORDA = Border(*[Side(style="thin", color="D9C7C4")] * 4)


def aba(wb, titulo, colunas, linhas, destacar=None):
    ws = wb.create_sheet(titulo) if wb.sheetnames != ["Sheet"] else wb.active
    if ws.title == "Sheet":
        ws.title = titulo
    ws.append([c[0] for c in colunas])
    for i, (_, larg) in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=i)
        cel.fill, cel.font = FILL_CAB, FONT_CAB
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    for linha in linhas:
        ws.append(list(linha))
    for r in range(2, ws.max_row + 1):
        marcar = destacar and destacar(ws, r)
        for c in range(1, len(colunas) + 1):
            cel = ws.cell(row=r, column=c)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
            cel.border = BORDA
            if marcar:
                cel.fill = FILL_BLOQ
            elif r % 2 == 0:
                cel.fill = FILL_ALT
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(colunas)), ws.max_row)
    return ws


def main():
    wb = Workbook()

    aba(wb, "Atribuicoes_Imprensa", COL_IMPRENSA, IMPRENSA,
        destacar=lambda ws, r: str(ws.cell(row=r, column=2).value or "").startswith("A COLETAR"))
    aba(wb, "Numeros_Documento_Legal", COL_LEGAL, LEGAL)
    aba(wb, "Leitura_Analitica", COL_ANALISE, ANALISE)
    aba(wb, "Protocolo_Coleta", COL_PROTO, PROTOCOLO)

    wb.save(SAIDA)
    print("{} — {} abas".format(SAIDA.name, len(wb.sheetnames)))
    for n in wb.sheetnames:
        print("   {:<26} {} linhas".format(n, wb[n].max_row - 1))
    bloq = sum(1 for r in IMPRENSA if r[1] == "A COLETAR")
    print("\n   verificados por leitura direta : {}".format(len(IMPRENSA) - bloq))
    print("   bloqueados, a coletar          : {}".format(bloq))


if __name__ == "__main__":
    main()
